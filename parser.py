import pdfplumber
import pandas as pd
import docx
import openpyxl
from fastapi import UploadFile
from io import BytesIO


def detect_file_type(filename: str):
    ext = filename.lower().split(".")[-1]
    return ext


# PDF parsing
def parse_pdf(file_bytes: bytes):
    data = []

    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if table:
                for row in table:
                    data.append(row)

    df = pd.DataFrame(data)
    return df


# Excel parsing
def parse_excel(file_bytes: bytes):
    workbook = openpyxl.load_workbook(BytesIO(file_bytes))
    sheet = workbook.active

    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(row)

    df = pd.DataFrame(rows)
    return df


# Word parsing
def parse_word(file_bytes: bytes):
    doc = docx.Document(BytesIO(file_bytes))

    rows = []
    for table in doc.tables:
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])

    df = pd.DataFrame(rows)
    return df

# CSV parsing
def parse_csv(file_bytes: bytes):
    df = pd.read_csv(BytesIO(file_bytes))
    return df


# Main parser entry point

async def parse_file(file: UploadFile):
    file_bytes = await file.read()
    ext = detect_file_type(file.filename)

    if ext == "pdf":
        return parse_pdf(file_bytes)

    elif ext in ["xlsx", "xls"]:
        return parse_excel(file_bytes)

    elif ext == "docx":
        return parse_word(file_bytes)

    elif ext == "csv":
        return parse_csv(file_bytes)

    else:
        raise ValueError(f"Unsupported file type: {ext}")
